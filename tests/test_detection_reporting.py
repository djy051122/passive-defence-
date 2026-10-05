from pathlib import Path

from openpyxl import load_workbook

from evaluation import (
    attach_ground_truth_labels,
    collect_sphere_parameters,
    count_four_class_predictions,
    decision_consistency_auc,
    evaluate_ground_truth_predictions,
    evaluate_sphere_purity,
    export_detection_workbook,
    format_final_report,
)


def _record(name: str, prediction: int, subtype: str = "N/A") -> dict:
    return {
        "path": str(Path("images") / name),
        "prediction": prediction,
        "prediction_name": ("real", "fake", "unknown")[prediction],
        "fake_subtype": subtype,
        "fake_subtype_confidence": 0.9 if prediction == 1 else None,
        "real_probability": 0.8 if prediction == 0 else 0.2,
        "fake_probability": 0.2 if prediction == 0 else 0.8,
        "confidence": 0.8,
        "novelty_score": 0.1,
        "sphere_distances": {"real": 0.4, "fake": 1.4},
        "sphere_memberships": {"real": prediction == 0, "fake": prediction == 1},
        "decision_region": (
            "real_sphere" if prediction == 0
            else "fake_sphere" if prediction == 1
            else "outside_all_spheres"
        ),
        "quality_reasons": [],
    }


def test_detection_console_summary_and_excel_export(tmp_path: Path):
    records = [
        _record("001_source.png", 0),
        _record("002_source.png", 2),
        _record("003_swapped.png", 1, "face_swap"),
        _record("F_STGN_004.jpg", 0),
    ]
    metrics = {
        "prediction_counts": {"real": 2, "fake": 1, "unknown": 1},
        "total_images": 4,
        "decision_region_counts": {
            "real_sphere": 2,
            "fake_sphere": 1,
            "outside_all_spheres": 1,
        },
    }
    artifacts = {
        "radii": {"real": 1.1, "fake": 1.2},
        "calibration_diagnostics": {
            "real": {
                "one_vs_rest_auc": 0.96,
                "purity": 0.95,
                "coverage": 0.90,
                "accepted_samples": 9,
            },
            "fake": {
                "one_vs_rest_auc": 0.94,
                "purity": 0.93,
                "coverage": 0.88,
                "accepted_samples": 8,
            },
        },
    }
    attach_ground_truth_labels(records)
    metrics["four_class_counts"] = count_four_class_predictions(records)
    metrics["ground_truth_evaluation"] = evaluate_ground_truth_predictions(records)
    metrics["sphere_purity"] = evaluate_sphere_purity(records)
    metrics["decision_consistency_auc"] = decision_consistency_auc(
        records, artifacts
    )
    metrics["sphere_parameters"] = collect_sphere_parameters(artifacts, metrics)

    report = format_final_report(metrics)
    assert "real sphere purity): 1/2 = 50.00%" in report
    assert "fake sphere purity): 1/1 = 100.00%" in report
    assert "真实球: 半径=1.1000, 参考AUC=0.9600" in report
    assert "模型决策一致性AUC=" in report
    assert "真实图片被判断为真: 1/2 = 50.00%" in report
    assert "假图片被判断为假: 1/2 = 50.00%" in report
    assert "换脸(face_swap): 1" in report
    assert "属性编辑(attribute_edit): 0" in report

    output = export_detection_workbook(
        records, metrics, artifacts, tmp_path / "detection_results.xlsx"
    )
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["检测汇总", "逐图结果"]
    summary = workbook["检测汇总"]
    assert summary["B7"].value == 4
    assert [summary[f"B{row}"].value for row in range(8, 12)] == [2, 1, 0, 1]
    assert summary["B14"].value == 1
    assert summary["C14"].value == 2
    assert summary["D14"].value == 0.5
    assert summary["B19"].value == 1.1
    assert summary["C19"].value == 0.96
    details = workbook["逐图结果"]
    assert details.max_row == 8
    assert [details[f"D{row}"].value for row in range(5, 9)] == [
        "真(real)", "真(real)", "假(fake)", "假(fake)",
    ]
    assert [details[f"F{row}"].value for row in range(5, 9)] == [
        "是", "否", "是", "否",
    ]
    assert [details[f"G{row}"].value for row in range(5, 9)] == [
        "真(real)",
        "未知(unknown)",
        "换脸(face_swap)",
        "真(real)",
    ]
