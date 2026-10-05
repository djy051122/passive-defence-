"""Console and Excel reporting helpers for final image detection."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from sklearn.metrics import roc_auc_score


FOUR_CLASS_NAMES = ("real", "face_swap", "attribute_edit", "unknown")
FOUR_CLASS_LABELS = {
    "real": "真(real)",
    "face_swap": "换脸(face_swap)",
    "attribute_edit": "属性编辑(attribute_edit)",
    "unknown": "未知(unknown)",
    "fake_unclassified": "未细分假图(fake_unclassified)",
}


def attach_ground_truth_labels(
    records: Sequence[dict[str, Any]],
    fake_subclass_names: Sequence[str] = ("face_swap", "attribute_edit"),
) -> None:
    """Attach independent labels after inference without feeding them to the model."""

    for record in records:
        known_label = int(record.get("known_label", -1))
        if known_label in (0, 1):
            label = known_label
            subtype = "N/A"
            subclass_index = int(record.get("fake_subclass", -1))
            if label == 1 and 0 <= subclass_index < len(fake_subclass_names):
                subtype = str(fake_subclass_names[subclass_index])
            source = "dataset_label"
        else:
            label, subtype = _ground_truth_from_filename(record.get("path", ""))
            source = "filename_rule" if label >= 0 else "unavailable"

        record["ground_truth_label"] = label
        record["ground_truth_name"] = (
            ("real", "fake")[label] if label in (0, 1) else "unlabeled"
        )
        record["ground_truth_fake_subtype"] = subtype
        record["ground_truth_source"] = source
        record["main_prediction_correct"] = (
            bool(int(record.get("prediction", 2)) == label)
            if label in (0, 1) else None
        )


def evaluate_ground_truth_predictions(
    records: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    """Calculate end-to-end real/fake recognition rates from independent labels."""

    real = [record for record in records if record.get("ground_truth_label") == 0]
    fake = [record for record in records if record.get("ground_truth_label") == 1]
    unlabeled = [
        record for record in records
        if int(record.get("ground_truth_label", -1)) not in (0, 1)
    ]
    real_correct = sum(int(record.get("prediction", 2)) == 0 for record in real)
    fake_correct = sum(int(record.get("prediction", 2)) == 1 for record in fake)
    return {
        "label_rule": (
            "*_source=real; *_swapped=face_swap; "
            "F_STGN_*=attribute_edit"
        ),
        "labeled_images": len(real) + len(fake),
        "unlabeled_images": len(unlabeled),
        "real_images": len(real),
        "real_predicted_real": real_correct,
        "real_correct_rate": real_correct / len(real) if real else None,
        "fake_images": len(fake),
        "fake_predicted_fake": fake_correct,
        "fake_correct_rate": fake_correct / len(fake) if fake else None,
    }


def evaluate_sphere_purity(
    records: Sequence[Mapping[str, Any]],
) -> dict[str, dict[str, Any]]:
    """Calculate geometric in-sphere purity from available ground-truth labels.

    Each sphere is evaluated independently from ``sphere_memberships``. Samples
    in an overlap therefore contribute to both spheres, while unlabeled samples
    are reported but excluded from the purity denominator.
    """

    result: dict[str, dict[str, Any]] = {}
    for name, expected_label in (("real", 0), ("fake", 1)):
        inside = [
            record for record in records
            if bool(record.get("sphere_memberships", {}).get(name, False))
        ]
        labeled_inside = [
            record for record in inside
            if int(record.get("ground_truth_label", -1)) in (0, 1)
        ]
        correct_inside = sum(
            int(record.get("ground_truth_label", -1)) == expected_label
            for record in labeled_inside
        )
        labeled_count = len(labeled_inside)
        result[name] = {
            "inside_samples": len(inside),
            "labeled_inside_samples": labeled_count,
            "unlabeled_inside_samples": len(inside) - labeled_count,
            "correct_inside_samples": correct_inside,
            "incorrect_inside_samples": labeled_count - correct_inside,
            "purity": correct_inside / labeled_count if labeled_count else None,
        }
    return result


def decision_consistency_auc(
    records: Sequence[Mapping[str, Any]],
    artifacts: Mapping[str, Any],
) -> dict[str, float | None]:
    """Measure agreement between final model decisions and sphere-distance scores.

    Predicted classes are deliberately used as pseudo-labels.  This is an
    internal consistency diagnostic, not a ground-truth performance AUC.
    """

    radii = artifacts.get("radii", {})
    result: dict[str, float | None] = {}
    for name, prediction_index in (("real", 0), ("fake", 1)):
        radius = _optional_float(radii.get(name))
        if radius is None or radius <= 0 or not records:
            result[name] = None
            continue
        labels = np.asarray([
            int(int(record.get("prediction", 2)) == prediction_index)
            for record in records
        ])
        if np.unique(labels).size < 2:
            result[name] = None
            continue
        scores = np.asarray([
            1.0 - float(record.get("sphere_distances", {}).get(name, np.inf)) / radius
            for record in records
        ])
        result[name] = float(roc_auc_score(labels, scores))
    return result


def _ground_truth_from_filename(path: str | Path) -> tuple[int, str]:
    stem = Path(str(path)).stem.lower()
    if stem.endswith("_source"):
        return 0, "N/A"
    if stem.endswith("_swapped"):
        return 1, "face_swap"
    if stem.startswith("f_stgn_"):
        return 1, "attribute_edit"
    return -1, "N/A"


def four_class_prediction(record: Mapping[str, Any]) -> str:
    """Map the three-output decision plus fake subtype to a display category."""

    prediction = int(record.get("prediction", 2))
    if prediction == 0:
        return "real"
    if prediction == 2:
        return "unknown"
    subtype = str(record.get("fake_subtype", ""))
    if subtype in ("face_swap", "attribute_edit"):
        return subtype
    return "fake_unclassified"


def count_four_class_predictions(
    records: Sequence[Mapping[str, Any]],
) -> dict[str, int]:
    counts = {name: 0 for name in FOUR_CLASS_NAMES}
    counts["fake_unclassified"] = 0
    for record in records:
        counts[four_class_prediction(record)] += 1
    return counts


def collect_sphere_parameters(
    artifacts: Mapping[str, Any], metrics: Mapping[str, Any]
) -> dict[str, dict[str, Any]]:
    """Collect radii and the most relevant available one-vs-rest AUC values."""

    radii = artifacts.get("radii", {})
    evaluated = metrics.get("spherewise_metrics", {})
    calibrated = artifacts.get("calibration_diagnostics", {})
    consistency = metrics.get("decision_consistency_auc", {})
    use_evaluation = all(
        isinstance(evaluated.get(name), Mapping) for name in ("real", "fake")
    )
    source = "当前有标签检测集" if use_evaluation else "校准集"
    result: dict[str, dict[str, Any]] = {}
    for name in ("real", "fake"):
        diagnostic = (
            evaluated.get(name, {}) if use_evaluation else calibrated.get(name, {})
        )
        result[name] = {
            "radius": _optional_float(radii.get(name)),
            "one_vs_rest_auc": _optional_float(
                diagnostic.get("one_vs_rest_auc")
            ),
            "auc_source": source,
            "decision_consistency_auc": _optional_float(consistency.get(name)),
            "purity": _optional_float(diagnostic.get("purity")),
            "class_coverage": _optional_float(
                diagnostic.get("class_coverage", diagnostic.get("coverage"))
            ),
            "accepted_samples": _optional_int(
                diagnostic.get("inside_samples", diagnostic.get("accepted_samples"))
            ),
        }
    return result


def export_detection_workbook(
    records: Sequence[Mapping[str, Any]],
    metrics: Mapping[str, Any],
    artifacts: Mapping[str, Any],
    output_path: str | Path,
) -> Path:
    """Write a readable two-sheet Excel report for one detection run."""

    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    counts = dict(
        metrics.get("four_class_counts") or count_four_class_predictions(records)
    )
    spheres = dict(
        metrics.get("sphere_parameters")
        or collect_sphere_parameters(artifacts, metrics)
    )

    workbook = Workbook()
    summary = workbook.active
    summary.title = "检测汇总"
    details = workbook.create_sheet("逐图结果")
    _write_summary_sheet(summary, len(records), counts, spheres, metrics)
    _write_details_sheet(details, records, spheres)
    workbook.save(target)
    return target


def _write_summary_sheet(sheet, total, counts, spheres, metrics) -> None:
    navy = "1F4E78"
    light_blue = "D9EAF7"
    thin_gray = Side(style="thin", color="D9E1F2")
    header_rows = (6, 13, 18, 23)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A7"
    sheet["A2"] = "深度伪造图片检测汇总"
    sheet["A2"].font = Font(name="Arial", size=14, bold=True, color="1F1F1F")
    sheet["A3"] = "AUC说明"
    sheet["B3"] = (
        "预测标签计算，仅表示模型内部一致性，不代表真实性能；参考AUC来自校准集或有标签测试集。"
    )
    sheet["A4"] = "真实标签规则"
    sheet["B4"] = "*_source=真；*_swapped=换脸；F_STGN_*=属性编辑；仅用于检测后核验。"
    for cell in (sheet["A3"], sheet["B3"], sheet["A4"], sheet["B4"]):
        cell.font = Font(name="Arial", size=10, italic=True, color="666666")

    sheet["A6"] = "模型预测类别"
    sheet["B6"] = "图片数量"
    category_rows = [
        ("总计", total),
        (FOUR_CLASS_LABELS["real"], counts.get("real", 0)),
        (FOUR_CLASS_LABELS["face_swap"], counts.get("face_swap", 0)),
        (FOUR_CLASS_LABELS["attribute_edit"], counts.get("attribute_edit", 0)),
        (FOUR_CLASS_LABELS["unknown"], counts.get("unknown", 0)),
    ]
    if counts.get("fake_unclassified", 0):
        category_rows.append(
            (FOUR_CLASS_LABELS["fake_unclassified"], counts["fake_unclassified"])
        )
    for row_index, row in enumerate(category_rows, start=7):
        sheet.cell(row_index, 1, row[0])
        sheet.cell(row_index, 2, row[1]).number_format = "#,##0"

    ground_truth = metrics.get("ground_truth_evaluation", {})
    ground_truth_headers = ["真实标签评价", "正确数量", "真实总数", "正确识别率"]
    for column, value in enumerate(ground_truth_headers, start=1):
        sheet.cell(13, column, value)
    ground_truth_rows = [
        (
            "真实图片被判断为真",
            ground_truth.get("real_predicted_real"),
            ground_truth.get("real_images"),
            ground_truth.get("real_correct_rate"),
        ),
        (
            "假图片被判断为假",
            ground_truth.get("fake_predicted_fake"),
            ground_truth.get("fake_images"),
            ground_truth.get("fake_correct_rate"),
        ),
    ]
    for row_index, values in enumerate(ground_truth_rows, start=14):
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value if value is not None else "n.a.")
        sheet.cell(row_index, 2).number_format = "#,##0"
        sheet.cell(row_index, 3).number_format = "#,##0"
        sheet.cell(row_index, 4).number_format = "0.00%"
    sheet.cell(16, 1, "无法匹配真实标签的图片")
    sheet.cell(16, 2, int(ground_truth.get("unlabeled_images", total)))
    sheet.cell(16, 2).number_format = "#,##0"

    sphere_headers = [
        "可信球", "半径", "参考AUC", "参考AUC来源", "模型决策一致性AUC",
        "纯度", "本类覆盖率", "球内样本数",
    ]
    for column, value in enumerate(sphere_headers, start=1):
        sheet.cell(18, column, value)
    for row_index, name in enumerate(("real", "fake"), start=19):
        item = spheres.get(name, {})
        values = [
            "真实球" if name == "real" else "伪造球",
            item.get("radius"),
            item.get("one_vs_rest_auc"),
            item.get("auc_source", "不可用"),
            item.get("decision_consistency_auc"),
            item.get("purity"),
            item.get("class_coverage"),
            item.get("accepted_samples"),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)
        for column in (2, 3, 5):
            sheet.cell(row_index, column).number_format = "0.0000"
        for column in (6, 7):
            sheet.cell(row_index, column).number_format = "0.00%"
        sheet.cell(row_index, 8).number_format = "#,##0"

    sheet.cell(23, 1, "判断区域")
    sheet.cell(23, 2, "图片数量")
    for offset, (region, value) in enumerate(
        sorted(metrics.get("decision_region_counts", {}).items()), start=1
    ):
        sheet.cell(23 + offset, 1, region)
        sheet.cell(23 + offset, 2, int(value))

    for header_row, last_column in ((6, 2), (13, 4), (18, 8), (23, 2)):
        for column in range(1, last_column + 1):
            cell = sheet.cell(header_row, column)
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin_gray)
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=8):
        for cell in row:
            if cell.row not in header_rows and cell.row != 2:
                cell.font = Font(
                    name="Arial", size=10,
                    italic=bool(cell.row in (3, 4)),
                    color="666666" if cell.row in (3, 4) else "1F1F1F",
                )
            cell.alignment = Alignment(vertical="center", wrap_text=False)
    for cell in (sheet["A7"], sheet["B7"]):
        cell.fill = PatternFill("solid", fgColor=light_blue)
    widths = {
        "A": 30, "B": 28, "C": 18, "D": 20,
        "E": 23, "F": 14, "G": 14, "H": 14,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[2].height = 24
    sheet.row_dimensions[3].height = 20
    sheet.row_dimensions[4].height = 20
    sheet["B3"].alignment = Alignment(vertical="center", wrap_text=False)
    sheet["B4"].alignment = Alignment(vertical="center", wrap_text=False)


def _write_details_sheet(sheet, records, spheres) -> None:
    navy = "1F4E78"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "D5"
    sheet["A2"] = "逐图检测结果"
    sheet["A2"].font = Font(name="Arial", size=14, bold=True, color="1F1F1F")
    headers = [
        "序号", "文件名", "完整路径", "真实主标签", "真实伪造类型", "主判断正确",
        "四分类结果", "主分类", "预测伪造类型", "子类型置信度", "真实概率",
        "伪造概率", "总体置信度", "新颖度分数", "到真实球距离", "真实球半径",
        "位于真实球", "到伪造球距离", "伪造球半径", "位于伪造球", "判断区域",
        "质量问题",
    ]
    header_row = 4
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, value)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    real_radius = spheres.get("real", {}).get("radius")
    fake_radius = spheres.get("fake", {}).get("radius")
    for row_index, record in enumerate(records, start=header_row + 1):
        category = four_class_prediction(record)
        distances = record.get("sphere_distances", {})
        memberships = record.get("sphere_memberships", {})
        quality_reasons = record.get("quality_reasons", [])
        values = [
            row_index - header_row,
            Path(str(record.get("path", ""))).name,
            str(record.get("path", "")),
            _ground_truth_label_text(record.get("ground_truth_name")),
            str(record.get("ground_truth_fake_subtype", "N/A")),
            _correctness_text(record.get("main_prediction_correct")),
            FOUR_CLASS_LABELS.get(category, category),
            str(record.get("prediction_name", "")),
            str(record.get("fake_subtype", "N/A")),
            record.get("fake_subtype_confidence"),
            record.get("real_probability"),
            record.get("fake_probability"),
            record.get("confidence"),
            record.get("novelty_score"),
            distances.get("real"),
            real_radius,
            "是" if memberships.get("real", False) else "否",
            distances.get("fake"),
            fake_radius,
            "是" if memberships.get("fake", False) else "否",
            str(record.get("decision_region", "")),
            "; ".join(str(item) for item in quality_reasons),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
        for column in list(range(10, 17)) + [18, 19]:
            sheet.cell(row_index, column).number_format = "0.0000"

    if records:
        last_row = header_row + len(records)
        table = Table(displayName="DetectionResultsTable", ref=f"A4:V{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    widths = [
        8, 24, 58, 14, 20, 14, 25, 12, 20, 16, 13,
        13, 13, 14, 17, 14, 14, 17, 14, 14, 24, 30,
    ]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[2].height = 24


def _optional_float(value: Any) -> float | None:
    return None if value is None else float(value)


def _optional_int(value: Any) -> int | None:
    return None if value is None else int(value)


def _correctness_text(value: Any) -> str:
    if value is None:
        return "无法评价"
    return "是" if bool(value) else "否"


def _ground_truth_label_text(value: Any) -> str:
    return {
        "real": "真(real)",
        "fake": "假(fake)",
        "unlabeled": "未标注",
    }.get(str(value), "未标注")
